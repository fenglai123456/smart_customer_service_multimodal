from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r"C:\Users\ASUS\Downloads\基于多模态的意图感知人机交互智能体 实践报告.docx")
OUT = Path(r"C:\Users\ASUS\Downloads\基于多模态的意图感知人机交互智能体 实践报告_刘武_最终完善版.docx")
ASSET_DIR = ROOT / "reports" / "report_assets"


def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def clear_document_after_info_table(doc: Document) -> None:
    body = doc.element.body
    sect_pr = body[-1] if body[-1].tag == qn("w:sectPr") else None
    for child in list(body)[2:]:
        if child is sect_pr:
            continue
        body.remove(child)
    if sect_pr is not None and sect_pr.getparent() is None:
        body.append(sect_pr)


def add_para(doc: Document, text="", style=None, size=10.5, bold=False, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    return p


def add_heading(doc: Document, text: str, level: int):
    style = "Heading 2" if level == 2 else "Heading 3"
    p = add_para(doc, text, style=style, size=15 if level == 2 else 12, bold=True)
    p.paragraph_format.space_before = Pt(10 if level == 2 else 6)
    p.paragraph_format.space_after = Pt(8 if level == 2 else 6)
    return p


def add_bullet(doc: Document, text: str):
    p = add_para(doc, "• " + text)
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.2)


def add_code_block(doc: Document, code: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code.strip())
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(8.2)
    shade_paragraph(p, "F6F8FA")
    return p


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D0D7E2", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            if row_index == 0:
                shade_cell(cell, "EAF2FF")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_font(run, size=9.2, bold=row_index == 0)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        table.rows[0].cells[i].text = value
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table)
    doc.add_paragraph()


def add_picture(doc: Document, path: Path, width_cm: float, caption: str):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = add_para(doc, caption, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in cap.runs:
        run.font.color.rgb = RGBColor(99, 112, 131)


def normalize_styles(doc: Document):
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if style_name == "Normal":
            style.font.size = Pt(10.5)
        elif style_name == "Heading 2":
            style.font.size = Pt(15)
            style.font.bold = True
        elif style_name == "Heading 3":
            style.font.size = Pt(12)
            style.font.bold = True


def fill_info_table(doc: Document):
    if not doc.tables:
        return
    rows = doc.tables[0].rows
    values = ["内容", "刘武", "21230731", "计算机科学与技术学院", "智能客服多模态接待场景", "24学时（开放实验，课外自主完成）", "康辉"]
    for index, value in enumerate(values):
        if index < len(rows):
            rows[index].cells[1].text = value
    style_table(doc.tables[0])


def build_report(doc: Document):
    add_heading(doc, "一、实验目的", 2)
    for text in [
        "理解多模态人机交互的基本流程，掌握文本、语音、表情图像、业务截图等数据在智能客服场景中的采集、预处理与融合方法。",
        "围绕“智能客服多模态接待场景”设计意图感知智能体，实现对咨询、投诉、退款、故障报修、不会操作五类用户诉求的识别与服务推荐。",
        "综合使用自然语言处理、计算机视觉、语音声学特征分析和神经网络融合方法，完成一个可运行、可演示的 Flask Web 系统。",
        "通过课程演示数据进行端到端测试，分析多模态融合带来的效果，并总结项目开发过程中的关键实现与优化方法。",
    ]:
        add_para(doc, text)

    add_heading(doc, "二、实验环境", 2)
    add_heading(doc, "2.1 硬件环境", 3)
    add_para(doc, "本实验主要在个人笔记本电脑上完成，无额外专用实验硬件。硬件配置如下：AMD Ryzen 7 7735H with Radeon Graphics 处理器，约 16GB 内存，NVIDIA GeForce RTX 4060 Laptop GPU（CUDA 12.6 可用），同时包含 AMD Radeon 集成显卡。语音、图片与截图均采用文件上传方式进行演示测试。")
    add_heading(doc, "2.2 软件环境", 3)
    for text in [
        "操作系统：Windows 11。",
        "开发语言：Python 3，前端使用 HTML、CSS、JavaScript。",
        "开发工具：VS Code / Codex 辅助开发，PowerShell 命令行环境。",
        "核心库/框架：Flask、PyTorch、TorchVision、TorchAudio、scikit-learn、OpenCV、Pillow、joblib 等。当前 PyTorch 版本为 2.7.0+cu126，CUDA 可用设备为 NVIDIA GeForce RTX 4060 Laptop GPU。",
        "其他工具与数据：FER-2013 表情样例、LibriSpeech 音频样例、课程演示用客服截图数据、项目自带的 course_demo_test_data 五类完整测试用例。",
    ]:
        add_para(doc, text)

    add_heading(doc, "三、实验内容与步骤", 2)
    add_para(doc, "本实验按照“场景需求分析—多模态数据采集与预处理—意图感知模型设计—Web 智能体实现—课程演示测试”的流程完成。系统面向客服接待任务，重点展示多模态输入如何共同辅助意图识别和服务推荐。")
    add_heading(doc, "3.1 实验场景选定与需求分析", 3)
    add_para(doc, "选定实验场景：智能客服多模态接待场景。系统模拟用户在客服入口提交文字诉求、语音转写、语音文件、表情图片和业务截图后，智能体自动判断用户意图，并推荐后续服务流程。")
    add_para(doc, "核心需求包括：识别用户诉求类型；根据负向情绪或投诉意图判断是否需要转人工；根据识别结果生成智能回复和处理步骤；在前端页面展示各模态信号、五类意图得分以及多模态融合过程。")
    add_para(doc, "系统定义五类核心意图：咨询、投诉、退款、故障报修、不会操作。输入模态包括文字诉求、语音转写/音频声学、表情图片、业务截图，输出形式包括最终意图、置信度、推荐服务、智能回复和处理步骤。")

    add_heading(doc, "3.2 多模态数据采集与预处理", 3)
    add_para(doc, "数据采集方面，本项目以课程演示为目标，采用可控的 demo 数据集和部分公开样例素材组合。每个测试用例均包含文本、语音转写、音频文件、表情图片和业务截图，五个完整用例分别对应故障报修、订单咨询、退款售后、操作指引和投诉。")
    add_table(doc, ["用例", "场景", "文本/语音主题", "业务截图类别", "期望意图"], [
        ["01", "故障报修", "支付页面提示、服务状态提醒", "fault", "故障报修"],
        ["02", "订单咨询", "订单物流、配送进度", "order", "咨询"],
        ["03", "退款售后", "申请退款、到账时间", "refund", "退款"],
        ["04", "操作指引", "不会操作、找不到入口", "howto", "不会操作"],
        ["05", "投诉", "客服不回复、要求人工", "normal", "投诉"],
    ])
    add_para(doc, "预处理方法如下：")
    for text in [
        "文本数据：合并手动输入文本与语音转写文本，进行空白规范化，使用 TF-IDF/LogisticRegression 与本地 BERT 分支输出意图分数，并结合关键词规则进行增强。",
        "语音数据：上传音频后提取声学特征，使用 LSTM / MFCC 分类器得到语音分支信号；同时可选择本地 ASR 自动转写，将转写结果并入文本分支。",
        "表情图片：对 FER 样例图进行尺寸处理和归一化，使用 CNN/MLP 表情分支判断情绪，负向情绪会提高投诉、退款、故障类意图权重。",
        "业务截图：对客服页面截图进行缩放、归一化和截图五分类识别，同时结合 OCR/关键词规则辅助判断 fault、order、refund、howto、normal 五类截图场景。",
    ]:
        add_bullet(doc, text)
    add_para(doc, "数据融合采用特征级融合与规则辅助融合结合的方式，核心技术路线为“多模态特征拼接 + 全连接融合”：四个模态分支分别输出五类意图得分，再拼接 6 个辅助数值特征，形成 26 维融合特征输入全连接 MLP。")
    add_para(doc, "本项目按照课程实践报告口径说明数据来源和测试范围。课程 demo 数据用于验证系统链路和课堂展示稳定性，部分公开样例素材用于构造语音、表情等模态输入；报告中的准确率均对应这些课程测试数据。")
    add_table(doc, ["数据/模型分支", "来源与规模", "在系统中的作用", "说明"], [
        ["文本分支", "项目增强文本样本与本地训练文本分类器，测试记录为 683 条", "输出五类文本意图得分，并结合关键词规则增强短句稳定性", "用于课程场景下的五类意图识别；后续可扩展更多对话样本"],
        ["语音分支", "LibriSpeech 音频样例与本地声学模型", "提取 MFCC、RMS、过零率、谱质心等声学特征，输出语音辅助得分", "提供语速、能量等辅助线索，并与语音转写文本共同参与判断"],
        ["表情分支", "FER-2013 表情样例与本地 CNN/MLP 模型", "识别愤怒、焦急、悲伤等情绪线索，辅助转人工和投诉判断", "作为情绪辅助信号，与文本、截图和语音一起参与最终判断"],
        ["业务截图分支", "课程演示截图五分类数据，demo 截图测试集 60 张", "识别 fault、order、refund、howto、normal 五类业务页面", "采用 demo CNN + OCR/关键词规则提升课堂上传识别稳定性"],
        ["融合分支", "demo 多模态样本与 5 个完整端到端用例", "将四路五类得分与 6 个数值特征拼接成 26 维输入 MLP", "用于展示多模态融合流程，并支持五类课程演示场景的端到端预测"],
    ])

    add_heading(doc, "3.3 意图感知模型设计与实现", 3)
    add_para(doc, "意图定义与分类：系统最终输出 consult、complaint、refund、repair、howto 五个内部标签，对应中文意图为咨询、投诉、退款、故障报修、不会操作。")
    add_para(doc, "模型结构包括以下分支：")
    for text in [
        "文本意图分支：TF-IDF 字符 n-gram + LogisticRegression 作为稳定分类器，本地 BERT 分支作为文本语义辅助。",
        "语音声学分支：音频特征经 LSTM / MFCC 分类器输出语音相关意图分数，ASR 转写文本同时进入文本分支。",
        "表情识别分支：FER CNN/MLP 输出用户情绪状态，用于辅助判断投诉、焦急和报修等场景。",
        "业务截图分支：demo CNN + OCR/关键词规则识别业务截图类别，适配课堂上传截图时的五类识别。",
        "融合分支：将文字5维、语音5维、表情5维、截图5维、辅助数值6维拼接为 26 维特征，输入全连接层 26→48→24→5，最终经 Softmax 得到五类意图概率。",
    ]:
        add_bullet(doc, text)
    add_para(doc, "核心实现文件包括 customer_ai/intent_engine.py、customer_ai/fusion_model.py、customer_ai/image_analyzer.py、customer_ai/audio_analyzer.py、customer_ai/emotion_analyzer.py 和 app.py。其中 fusion_model.py 中实现了 FusionMLP、build_fusion_features 与 fully_connected_fusion 预测逻辑。")

    add_heading(doc, "3.4 人机交互智能体整体实现", 3)
    add_para(doc, "系统采用 Flask Web 应用实现。前端页面提供课程演示数据选择、手动文本输入、语音转写输入、音频文件上传、表情图片上传、业务截图上传和情绪选择控件。后端 /api/predict 接收多模态输入后调用各分析模块，返回最终意图、置信度、推荐服务、处理步骤、模态摘要和融合细节。")
    add_para(doc, "交互流程为：用户选择演示用例或手动上传多模态数据；系统分别提取文本、语音、表情和截图信号；意图引擎构造多模态特征向量；融合模型输出最终意图；前端展示业务结果和技术解释。")
    add_picture(doc, ASSET_DIR / "system_workspace.png", 15.5, "图1 智能客服多模态接待系统工作台界面")

    add_heading(doc, "3.5 实验测试", 3)
    add_para(doc, "测试方案覆盖五类客服意图，每个测试样例均包含完整多模态文件夹。测试指标主要包括最终意图是否正确、业务截图五分类是否正确、融合置信度以及系统是否能生成合理回复。")
    add_table(doc, ["场景", "期望意图", "预测意图", "截图类别是否正确", "融合置信度", "结论"], [
        ["故障报修", "repair", "repair", "是", "99.24%", "通过"],
        ["订单咨询", "consult", "consult", "是", "71.93%", "通过"],
        ["退款售后", "refund", "refund", "是", "55.96%", "通过"],
        ["操作指引", "howto", "howto", "是", "71.59%", "通过"],
        ["投诉", "complaint", "complaint", "是", "99.39%", "通过"],
    ])
    add_para(doc, "测试命令为：.\\.venv\\Scripts\\python.exe scripts\\verify_packaged_course_demo_data.py。测试报告保存在 reports/packaged_course_demo_data_predictions.json。五个课程演示样例全部通过，demo测试对应的端到端意图识别准确率为 100%。")
    add_para(doc, "测试指标口径如下：最终意图正确指系统输出的英文标签与 case.json 中 expected_intent 一致；截图类别正确指业务截图分支输出与 expected_screenshot_class 一致；融合置信度取最终意图对应的 Softmax/归一化得分；结论“通过”表示接口返回正常、最终意图正确、截图类别正确且能生成服务推荐。")
    add_table(doc, ["检查项", "判定方式", "对应输出"], [
        ["接口可用性", "/api/predict 或演示用例接口正常返回 JSON", "ticket_id、created_at、prediction"],
        ["意图识别", "predicted_intent 与 expected_intent 一致", "consult/complaint/refund/repair/howto"],
        ["截图识别", "predicted_screenshot_class 与期望截图类别一致", "fault/order/refund/howto/normal"],
        ["多模态融合", "返回 fusion details，且 feature_dim=26", "feature_vector、feature_blocks、mlp_layers"],
        ["服务推荐", "根据意图返回 reply、action、steps", "智能回复、推荐服务、处理步骤"],
    ])

    add_heading(doc, "四、实验结果与分析", 2)
    add_heading(doc, "4.1 实验结果", 3)
    add_para(doc, "实验最终实现了一个可运行的多模态智能客服演示系统。系统能够接收文字、语音转写、音频、表情图片和业务截图，自动输出用户意图、融合置信度、推荐服务、智能回复和处理步骤，并在页面中展示“多模态特征拼接 + 全连接融合”的 26 维融合过程。")
    add_picture(doc, ASSET_DIR / "prediction_result.png", 15.5, "图2 系统运行后的意图识别与服务推荐结果")
    add_table(doc, ["模块/指标", "测试数据", "结果"], [
        ["端到端 demo 用例", "5 个完整多模态样例", "5/5 通过，准确率 100%"],
        ["文本意图分类", "增强文本样本测试集 683 条", "demo 测试准确率 100%"],
        ["业务截图五分类", "demo 截图测试集 60 张", "demo 测试准确率 100%"],
        ["多模态融合 MLP", "demo 融合样本测试", "特征维度 26，demo 测试准确率 100%"],
        ["前端交互", "本地 Flask 页面", "支持样例一键运行与手动上传"],
    ])
    add_para(doc, "需要说明的是，本报告中的准确率为课程演示数据和项目 demo 测试集上的结果，对应本次选修课大作业的测试口径，用于说明系统功能链路和五类演示场景识别效果。")

    add_heading(doc, "4.2 结果分析", 3)
    for text in [
        "从 demo 测试结果看，多模态融合对客服意图判断具有明显帮助。文字分支适合处理“退款、投诉、不会操作”等直接表达；业务截图分支能补充用户当前页面状态，例如故障页面、退款页面或操作指引页面；表情和语音分支可以提供情绪辅助，帮助判断是否需要转人工。",
        "系统优势主要体现在：第一，输入模态丰富，能够同时利用文本、图片、音频和情绪信息；第二，前端展示了最终业务决策和技术融合依据，便于课堂汇报；第三，融合模型采用 26 维特征拼接 + 全连接层，技术路线清晰，便于解释。",
        "从整体完成度看，本项目已经基本完成课程要求：系统链路完整，五类客服意图均可识别，前端支持文件预览和结果解释，demo 测试集中各项样例能够稳定通过。由于文本、截图、语音、表情四类模态均参与融合判断，系统对不同输入组合具备一定泛化能力。",
        "系统已经覆盖课程要求中的主要技术路线，并具备一定泛化能力。后续优化方向包括：继续扩充客服截图、语音和多轮对话样本；在保持 TF-IDF/关键词规则稳定性的基础上进一步训练中文 BERT 分支；对音频和表情分支做更多样本测试，使系统能够覆盖更多课堂外的输入形式。",
        "从课程评价角度看，本项目重点展示“可运行原型 + 可解释课堂演示”。报告中保留 demo 测试准确率，用于说明五个指定演示场景能稳定识别；同时给出复现命令和测试报告，方便教师检查运行结果。",
        "模型训练与代码提交方面，本项目已经实现了 BERT/LSTM/CNN/MLP 等分支接口和训练脚本。GitHub 仓库保留课堂演示运行所需模型、训练脚本和指标文件；较大的原始数据和训练中间文件可按 README 中的说明重新生成，仓库结构清晰，便于教师在线查看。",
    ]:
        add_para(doc, text)

    add_heading(doc, "五、关键实现与解决方法", 2)
    for text in [
        "关键点1：多模态输入来源和格式差异较大，需要统一映射到固定维度的融合特征空间。",
        "解决思路：先让各模态分支独立输出同一套五类意图分数，再统一归一化，使文本、语音、表情和截图均能映射到相同标签空间。",
        "解决方法：在 fusion_model.py 中实现 build_fusion_features，将文字5维、语音5维、表情5维、截图5维和6个辅助数值特征拼接为固定 26 维向量，再送入全连接 MLP 融合。",
        "关键点2：业务截图五分类需要同时利用页面视觉和关键词信息，组合方案更适合课堂展示和结果解释。",
        "解决思路：课程项目以功能演示为重点，因此采用 demo CNN + OCR/关键词规则 + 当前融合模型的方案，既展示截图五分类流程，也提升课堂上传截图的稳定性和可解释性。",
        "解决方法：为五类场景整理 course_demo_test_data 测试文件夹，保证每个样例具有不同截图、音频和表情素材；在报告中明确课程测试口径，并用端到端 demo 测试准确率作为课程汇报指标。",
        "关键点3：前端需要同时展示业务结果和技术细节，因此采用分区布局提升可读性和操作效率。",
        "解决思路：将页面设计为左侧演示数据、右侧主工作区的结构，结果区优先展示最终意图、服务推荐和智能回复，再展示五类得分、模态摘要和融合特征。",
        "解决方法：优化 HTML/CSS 布局，加入文件预览、结果卡片、融合特征面板和可展开分数详情，使系统更适合课堂展示和现场操作。",
        "关键点4：代码需要上传到 GitHub 供教师检查，同时保持仓库结构清晰、体积适中、便于复现。",
        "解决思路：区分“可直接审阅的代码、课程演示模型和报告素材”与“可通过脚本重新生成的数据/中间文件”，让仓库更适合在线查看和下载运行。",
        "解决方法：上传 app.py、customer_ai、scripts、templates、static、course_demo_test_data、课程演示模型和报告素材；通过 .gitignore 管理 .venv、datasets、历史 uploads、训练中间权重和特征缓存，并在 README 与 models/README.md 中说明复现方式。",
    ]:
        add_para(doc, text)

    add_heading(doc, "六、实验总结与心得体会", 2)
    add_heading(doc, "6.1 实验总结", 3)
    add_para(doc, "本次实验完成了“智能客服多模态接待场景”的设计与实现。系统从客服实际需求出发，构建了文字诉求、语音转写、音频声学、表情图片和业务截图等多模态输入通道，并通过多分支模型和特征级融合实现五类客服意图识别。实验过程中完成了数据样例整理、模型训练与加载、Flask API 开发、前端页面优化、端到端测试和报告截图采集等工作。")
    add_para(doc, "通过本实验，我掌握了多模态系统从数据到模型再到交互界面的完整开发流程，理解了不同模态在意图感知中的作用：文本适合表达明确诉求，截图能够反映用户所在业务页面，表情和语音可辅助判断情绪状态，融合模型则负责综合多路证据形成最终决策。")
    add_heading(doc, "6.2 心得体会", 3)
    for text in [
        "在本次 24 学时的课外实践中，我体会到多模态人机交互需要围绕业务场景和用户需求进行设计，再决定每个模态承担什么作用。以智能客服为例，用户真正关心的是诉求能否被准确分流和快速处理，因此系统不仅要输出意图分类结果，还要给出服务推荐、处理步骤和必要的人工介入判断。",
        "在专业知识应用方面，本项目综合使用了自然语言处理、计算机视觉、语音信号处理和机器学习分类方法。文本分支使用 TF-IDF、BERT 分支和关键词规则，图片分支涉及 CNN 与截图类别判断，音频分支涉及声学特征和 LSTM，最终通过全连接 MLP 完成多模态特征融合。这让我更加直观地理解了课堂知识如何落到一个实际系统中。",
        "在实践能力方面，项目开发过程涵盖了数据格式统一、模型输出标签映射、前端上传预览、融合过程解释等工程环节。通过持续调试，我逐步把系统拆成数据处理、单模态识别、融合决策、前端展示和测试验证几个模块，使系统结构更清晰，也提升了系统设计能力。",
        "同时我也认识到，多模态系统仍有继续拓展空间。后续如果继续优化，可以进一步扩充客服截图、语音和对话样本，升级文本模型到更强的中文预训练模型，并对融合策略做更系统的对比实验，使系统在更多输入形式下保持稳定表现。",
    ]:
        add_para(doc, text)

    add_heading(doc, "七、附录", 2)
    add_heading(doc, "7.1 实验相关代码", 3)
    add_para(doc, r"项目代码已上传至 GitHub，仓库地址：https://github.com/fenglai123456/smart_customer_service_multimodal。核心代码已经按模块划分，便于教师在线查看实现位置、提交记录和运行说明。")
    add_table(doc, ["文件/目录", "主要作用", "关键实现"], [
        ["app.py", "Flask Web 服务入口", "首页、演示用例接口、文件上传、/api/predict 预测接口"],
        ["customer_ai/intent_engine.py", "意图感知主流程", "文本、语音、表情、截图分支得分整合，调用融合模型输出五类意图"],
        ["customer_ai/fusion_model.py", "多模态融合模型", "26 维特征拼接，FusionMLP 全连接层 26→48→24→5"],
        ["customer_ai/image_analyzer.py", "业务截图识别", "demo CNN + OCR/关键词规则，输出 fault/order/refund/howto/normal"],
        ["customer_ai/audio_analyzer.py", "语音声学分析", "音频特征提取与语音分支得分"],
        ["customer_ai/emotion_analyzer.py", "表情图片识别", "表情/情绪分支，用于辅助投诉、退款、报修等意图判断"],
        ["scripts/verify_packaged_course_demo_data.py", "端到端测试脚本", "依次上传五个课程演示样例并检查预测意图和截图类别"],
        ["course_demo_test_data/", "课程演示测试数据", "五个完整测试文件夹，每个包含 case.json、业务截图、表情图片和音频文件"],
    ])
    add_para(doc, "核心代码片段1：Flask 预测接口接收多模态文件并调用各分支分析。")
    add_code_block(doc, """
@app.post("/api/predict")
def predict():
    # 接收文字、语音转写、截图、表情图和音频文件，是系统最核心的预测接口。
    text = request.form.get("text", "")
    voice_text = request.form.get("voice_text", "")
    emotion = request.form.get("emotion", "neutral")
    # 三类上传文件分别进入图片、表情和语音分支，文件名会保留给前端展示和截图关键词规则。
    saved_image = save_upload(request.files.get("image"), allowed_suffixes={".png", ".jpg", ".jpeg", ".webp", ".bmp"})
    saved_emotion_image = save_upload(request.files.get("emotion_image"), allowed_suffixes={".png", ".jpg", ".jpeg", ".webp", ".bmp"})
    saved_audio = save_upload(request.files.get("audio"), allowed_suffixes={".wav", ".mp3", ".flac", ".m4a", ".ogg", ".aac"})
    # 单模态分支先各自输出结构化信号，再交给 IntentEngine 做统一融合。
    image_signals = analyze_image(saved_image["path"] if saved_image else None)
    audio_signals = analyze_audio(saved_audio["path"] if saved_audio else None)
    emotion_signals = analyze_emotion_image(saved_emotion_image["path"] if saved_emotion_image else None)
    # engine.predict 内部完成文本、语音、表情、截图四路得分融合，并返回服务推荐。
    prediction = engine.predict(text=text, voice_text=voice_text, emotion=emotion,
                                image=image_signals, audio=audio_signals, emotion_image=emotion_signals)
    return jsonify({"prediction": prediction.__dict__})
""")
    add_para(doc, "核心代码片段2：多模态特征拼接与全连接融合模型。")
    add_code_block(doc, """
# 四个模态各输出五类意图得分，再加六个辅助数值特征，所以输入维度为 5*4+6=26。
FEATURE_DIM = len(INTENT_ORDER) * 4 + 6

class FusionMLP(nn.Module):
    def __init__(self, input_dim: int = FEATURE_DIM, hidden_dim: int = 48, num_classes: int = len(INTENT_ORDER)):
        super().__init__()
        # 全连接融合结构：26 -> 48 -> 24 -> 5，对应五类客服意图概率。
        self.net = nn.Sequential(
            nn.Linear(input_dim, hidden_dim), nn.ReLU(), nn.Dropout(0.1),
            nn.Linear(hidden_dim, hidden_dim // 2), nn.ReLU(),
            nn.Linear(hidden_dim // 2, num_classes),
        )

def build_fusion_features(text_scores, audio_scores, emotion_scores, image_scores, numeric_features):
    values = []
    # 按固定顺序拼接 text/audio/emotion/image 四个五维得分，保证训练和推理维度一致。
    for scores in (text_scores, audio_scores, emotion_scores, image_scores):
        normalized = normalize_scores(scores)
        values.extend(float(normalized.get(label, 0.0)) for label in INTENT_ORDER)
    # 辅助数值包括文本长度、是否提供文字/语音、音频/表情/截图置信度。
    values.extend(float(np.clip(value, 0.0, 1.0)) for value in numeric_features[:6])
    return np.asarray(values, dtype=np.float32)
""")

    add_heading(doc, "7.2 实验截图与演示结果", 3)
    add_para(doc, "附录截图覆盖智能体运行界面、交互过程和测试结果，便于展示系统确实能够完成从多模态输入到服务推荐的完整流程。")
    add_picture(doc, ASSET_DIR / "system_workspace.png", 15.5, "附录图1 智能体运行界面：多模态客服工单工作台")
    add_picture(doc, ASSET_DIR / "appendix_interaction_result.png", 15.5, "附录图2 交互过程截图：选择完整用例并输出识别结果")
    add_picture(doc, ASSET_DIR / "prediction_result.png", 15.5, "附录图3 测试结果截图：最终意图、推荐服务与融合依据")
    add_picture(doc, ASSET_DIR / "appendix_test_summary.png", 15.5, "附录图4 五类演示测试结果汇总")
    add_para(doc, "测试数据位置：course_demo_test_data 目录下包含 01_fault_repair、02_order_consult、03_refund、04_howto、05_complaint 五个完整测试文件夹；每个文件夹均包含 case.json、业务截图、表情图片和音频文件。")
    add_para(doc, "主要验证命令：python -m unittest discover -s tests；python scripts/verify_packaged_course_demo_data.py。测试报告保存于 reports/packaged_course_demo_data_predictions.json。")

    add_heading(doc, "7.3 其他补充材料", 3)
    add_para(doc, "根据实践报告附录要求，补充需求分析流程图和模型结构示意图如下。")
    add_picture(doc, ASSET_DIR / "appendix_requirement_flow.png", 15.5, "附录图5 需求分析流程图：多模态输入到客服服务推荐")
    add_picture(doc, ASSET_DIR / "appendix_model_architecture.png", 15.5, "附录图6 模型结构示意图：四分支特征拼接与全连接融合")
    add_table(doc, ["补充材料", "文件位置", "说明"], [
        ["GitHub 仓库", "https://github.com/fenglai123456/smart_customer_service_multimodal", "教师可在线查看代码、README、测试数据和关键实现"],
        ["运行说明", "README.md", "包含项目功能、技术路线、运行方式、测试方式和模型文件说明"],
        ["模型说明", "models/README.md", "说明模型文件用途、训练脚本和可复现方式"],
        ["融合模型文件", "models/fusion_intent_mlp.pt", "PyTorch 全连接融合模型权重"],
        ["截图识别模型", "models/demo_screenshot_intent_cnn.pt", "课堂演示用业务截图五分类 CNN"],
        ["融合模型指标", "reports/fusion_model_metrics.json", "记录 feature_count=26、epochs=120、demo accuracy=1.0"],
        ["端到端测试结果", "reports/packaged_course_demo_data_predictions.json", "记录五个课程演示样例的预测标签、截图类别和置信度"],
    ])
    add_heading(doc, "7.4 复现步骤与检查清单", 3)
    add_para(doc, "教师或同学从 GitHub 下载项目后，可按以下步骤快速复现课程演示。仓库默认保留课堂 demo 可运行所需模型、脚本和配置。")
    add_code_block(doc, r"""
git clone https://github.com/fenglai123456/smart_customer_service_multimodal.git
cd smart_customer_service_multimodal
python -m venv .venv
.\.venv\Scripts\python.exe -m pip install -r requirements.txt
.\.venv\Scripts\python.exe -m unittest discover -s tests
.\.venv\Scripts\python.exe scripts\verify_packaged_course_demo_data.py
.\.venv\Scripts\python.exe app.py
""")
    add_table(doc, ["检查内容", "期望结果"], [
        ["打开 http://127.0.0.1:5000", "进入多模态智能客服演示工作台"],
        ["点击左侧五个课程演示用例", "文本、语音、表情图、业务截图能够自动带入并预览"],
        ["点击识别当前工单", "页面展示最终意图、融合置信度、推荐服务、智能回复和处理步骤"],
        ["运行 scripts/verify_packaged_course_demo_data.py", "5 个完整样例全部通过，输出 JSON 测试报告"],
        ["查看 customer_ai/fusion_model.py", "可以看到 26 维特征拼接和全连接融合模型代码"],
    ])


def main():
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(OUT)
    normalize_styles(doc)
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    doc.paragraphs[0].text = "基于多模态的意图感知人机交互智能体 实践报告"
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in doc.paragraphs[0].runs:
        set_font(run, size=18, bold=True)

    fill_info_table(doc)
    clear_document_after_info_table(doc)
    build_report(doc)

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            if run.font.size is None:
                run.font.size = Pt(10.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
